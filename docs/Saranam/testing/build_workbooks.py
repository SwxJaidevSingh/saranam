#!/usr/bin/env python3
"""
Build the Saranam India Health Care manual-testing Excel pack.

Reads every <module>/<module>_test_cases.json under docs/Saranam/testing/ and produces:
  - xlsx/<NN>_<Module>_Test_Cases.xlsx  (one fillable workbook per module)
  - xlsx/00_Saranam_Master_Test_Plan.xlsx (all modules as sheets + a live Dashboard)

The testing team fills the yellow columns: Actual Result, Status, Severity, Tested By,
Date Tested, Remarks / Defect ID. Status & Severity are dropdowns; Status is colour-coded.

Run:  python3 docs/Saranam/testing/build_workbooks.py
"""
import json
import glob
import os
import math

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "xlsx")
os.makedirs(OUT, exist_ok=True)

# ---- palette -------------------------------------------------------------
NAVY      = "1F3864"
HEADER_BG = "1F3864"
HEADER_FG = "FFFFFF"
FILL_BG   = "FFF2CC"   # tester-fill columns (light amber)
SECTION_BG= "D9E1F2"
ZEBRA     = "F2F6FC"
PASS_BG   = "C6EFCE"; PASS_FG = "006100"
FAIL_BG   = "FFC7CE"; FAIL_FG = "9C0006"
BLOCK_BG  = "FFEB9C"; BLOCK_FG= "9C6500"
NR_BG     = "E7E6E6"; NR_FG   = "595959"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

STATUS_OPTS   = ["Not Executed", "Pass", "Fail", "Blocked", "N/A"]
SEVERITY_OPTS = ["", "Critical", "High", "Medium", "Low"]

# Test-case table columns: (header, json_key_or_None, width, wrap, tester_fill)
COLS = [
    ("Test Case ID",        "id",              13, False, False),
    ("Section / Feature",   "section",         22, True,  False),
    ("Test Title",          "title",           34, True,  False),
    ("Priority",            "priority",         9, False, False),
    ("Type",                "type",            13, True,  False),
    ("Preconditions",       "preconditions",   26, True,  False),
    ("Test Steps",          "steps",           50, True,  False),
    ("Test Data",           "test_data",       24, True,  False),
    ("Expected Result",     "expected_result", 44, True,  False),
    ("API Reference",       "api_reference",   30, True,  False),
    ("Actual Result",       None,              30, True,  True),
    ("Status",              None,              14, False, True),
    ("Severity",            None,              11, False, True),
    ("Tested By",           None,              14, False, True),
    ("Date Tested",         None,              13, False, True),
    ("Remarks / Defect ID", None,              30, True,  True),
]
STATUS_COL = next(i for i, c in enumerate(COLS) if c[0] == "Status") + 1   # 1-based
SEV_COL    = next(i for i, c in enumerate(COLS) if c[0] == "Severity") + 1
STATUS_LTR = get_column_letter(STATUS_COL)
SEV_LTR    = get_column_letter(SEV_COL)


def steps_to_text(steps):
    if isinstance(steps, list):
        return "\n".join(f"{i}. {s}" for i, s in enumerate(steps, 1))
    return str(steps or "")


def est_row_height(values):
    """Rough auto-height so wrapped cells are readable on open."""
    lines = 1
    for (header, key, width, wrap, _), val in zip(COLS, values):
        if not wrap or not val:
            continue
        cell_lines = 0
        for ln in str(val).split("\n"):
            cell_lines += max(1, math.ceil(len(ln) / max(8, width - 2)))
        lines = max(lines, cell_lines)
    return min(409, max(30, lines * 15))


def write_test_sheet(ws, data):
    """Write the standard Test Cases sheet onto worksheet `ws`."""
    tcs = data["test_cases"]
    # header
    for ci, (header, key, width, wrap, fill) in enumerate(COLS, 1):
        c = ws.cell(row=1, column=ci, value=header)
        c.font = Font(bold=True, color=HEADER_FG, size=10)
        c.fill = PatternFill("solid", fgColor=HEADER_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 30

    prev_section = None
    band = False
    for ri, tc in enumerate(tcs, start=2):
        if tc.get("section") != prev_section:
            band = not band
            prev_section = tc.get("section")
        row_vals = []
        for header, key, width, wrap, fill in COLS:
            if key == "steps":
                v = steps_to_text(tc.get("steps"))
            elif key is None:
                v = "Not Executed" if header == "Status" else ""
            else:
                v = tc.get(key, "")
            row_vals.append(v)

        for ci, ((header, key, width, wrap, fill), v) in enumerate(zip(COLS, row_vals), 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.alignment = Alignment(
                horizontal="left" if wrap else "center",
                vertical="top", wrap_text=wrap)
            c.border = BORDER
            c.font = Font(size=9)
            if fill:
                c.fill = PatternFill("solid", fgColor=FILL_BG)
            elif band:
                c.fill = PatternFill("solid", fgColor=ZEBRA)
            if header == "Test Case ID":
                c.font = Font(size=9, bold=True)
            if header == "Priority":
                pf = {"High": "9C0006", "Medium": "9C6500", "Low": "375623"}.get(str(v), "000000")
                c.font = Font(size=9, bold=True, color=pf)
        ws.row_dimensions[ri].height = est_row_height(row_vals)

    last = len(tcs) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{last}"

    if tcs:
        dv_status = DataValidation(type="list", formula1='"%s"' % ",".join(STATUS_OPTS), allow_blank=True)
        dv_sev = DataValidation(type="list", formula1='"%s"' % ",".join([s for s in SEVERITY_OPTS if s]), allow_blank=True)
        ws.add_data_validation(dv_status)
        ws.add_data_validation(dv_sev)
        dv_status.add(f"{STATUS_LTR}2:{STATUS_LTR}{last}")
        dv_sev.add(f"{SEV_LTR}2:{SEV_LTR}{last}")

        rng = f"{STATUS_LTR}2:{STATUS_LTR}{last}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Pass"'],
            fill=PatternFill("solid", fgColor=PASS_BG), font=Font(color=PASS_FG, bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fail"'],
            fill=PatternFill("solid", fgColor=FAIL_BG), font=Font(color=FAIL_FG, bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Blocked"'],
            fill=PatternFill("solid", fgColor=BLOCK_BG), font=Font(color=BLOCK_FG, bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Not Executed"'],
            fill=PatternFill("solid", fgColor=NR_BG), font=Font(color=NR_FG)))
    return last


def write_info_sheet(ws, data, last_row_in_tc):
    """Cover / scope / summary sheet for a single-module workbook."""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 95
    r = 1

    def title(text):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(bold=True, size=14, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 26
        r += 1

    def kv(k, v):
        nonlocal r
        a = ws.cell(row=r, column=1, value=k); a.font = Font(bold=True, size=10)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        a.fill = PatternFill("solid", fgColor=SECTION_BG)
        b = ws.cell(row=r, column=2, value=v); b.font = Font(size=10)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = max(18, min(160, 15 * (1 + str(v).count("\n") + len(str(v)) // 95)))
        r += 1

    def gap():
        nonlocal r
        r += 1

    title(f"{data['module']}  —  Manual Test Plan")
    kv("Platform", "Saranam India Health Care  (https://saranam.fieldintel.ai)")
    kv("Module Code", data.get("module_code", ""))
    kv("Document Version", data.get("version", "1.0"))
    kv("Total Test Cases", str(len(data["test_cases"])))
    kv("Test Type", "Manual")
    kv("Prepared / Reviewed", "____________________     Date: ____________")
    gap()

    title("1. Scope")
    kv("Summary", data.get("summary", ""))
    kv("In Scope", "\n".join("• " + s for s in data.get("in_scope", [])))
    kv("Out of Scope", "\n".join("• " + s for s in data.get("out_of_scope", [])) or "—")
    gap()

    title("2. Roles & Access")
    kv("Access", data.get("roles_and_access", ""))
    gap()

    title("3. Screens & Routes")
    for s in data.get("screens", []):
        kv(s.get("name", ""), s.get("route", ""))
    gap()

    title("4. API Endpoints")
    eps = data.get("api_endpoints", [])
    if eps:
        for e in eps:
            kv(f"{e.get('method','')}  {e.get('path','')}", e.get("purpose", ""))
    else:
        kv("—", "None")
    gap()

    title("5. Test Data / Preconditions")
    kv("Setup", "\n".join("• " + s for s in data.get("test_data_setup", [])) or "—")
    gap()

    title("6. Execution Summary  (auto-calculated)")
    rng = f"'Test Cases'!{STATUS_LTR}2:{STATUS_LTR}{last_row_in_tc}"
    metrics = [
        ("Total Cases", f'=COUNTA(\'Test Cases\'!A2:A{last_row_in_tc})'),
        ("Pass",         f'=COUNTIF({rng},"Pass")'),
        ("Fail",         f'=COUNTIF({rng},"Fail")'),
        ("Blocked",      f'=COUNTIF({rng},"Blocked")'),
        ("N/A",          f'=COUNTIF({rng},"N/A")'),
        ("Not Executed", f'=COUNTIF({rng},"Not Executed")'),
        ("% Executed",   f'=IF(B{r}=0,0,(B{r}-COUNTIF({rng},"Not Executed"))/B{r})'),
        ("% Pass (of executed)", f'=IFERROR(B{r+1}/(B{r}-B{r+5}),0)'),
    ]
    base = r
    for i, (k, f) in enumerate(metrics):
        a = ws.cell(row=base + i, column=1, value=k); a.font = Font(bold=True, size=10)
        a.fill = PatternFill("solid", fgColor=SECTION_BG)
        b = ws.cell(row=base + i, column=2, value=f); b.font = Font(size=10)
        if k.startswith("%"):
            b.number_format = "0.0%"
    r = base + len(metrics) + 1

    title("7. Defect Severity Legend")
    for k, v in [
        ("Critical", "Blocks core workflow / data loss / security breach; no workaround. Fix before go-live."),
        ("High", "Major feature broken or incorrect; workaround difficult. Fix before go-live."),
        ("Medium", "Feature works with a usable workaround or minor data issue."),
        ("Low", "Cosmetic / minor UI / wording. Backlog."),
    ]:
        kv(k, v)
    gap()
    title("8. Status Values")
    kv("Pass / Fail / Blocked / N/A / Not Executed",
       "Pass = behaves as expected. Fail = does not match Expected Result (log Severity + Remarks/Defect ID). "
       "Blocked = cannot test due to a dependency/defect. N/A = not applicable in this environment. "
       "Not Executed = not yet run (default).")


def build_module_workbook(data, fname):
    wb = Workbook()
    tc = wb.active
    tc.title = "Test Cases"
    last = write_test_sheet(tc, data)
    info = wb.create_sheet("Module Info")
    write_info_sheet(info, data, last)
    # order: Module Info first
    wb.move_sheet("Module Info", offset=-1)
    wb.active = wb["Module Info"]
    path = os.path.join(OUT, fname)
    wb.save(path)
    return path, last


def build_master(modules):
    """modules: list of dicts {code,name,sheet,data,fname}"""
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    dash.sheet_view.showGridLines = False

    # title
    dash.merge_cells("A1:J1")
    t = dash.cell(row=1, column=1, value="Saranam India Health Care — Master Manual Test Plan")
    t.font = Font(bold=True, size=16, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    dash.row_dimensions[1].height = 30
    dash.merge_cells("A2:J2")
    s = dash.cell(row=2, column=1, value="Manual test execution tracker • Fill Status on each module sheet • figures below update automatically")
    s.font = Font(italic=True, size=10, color="595959")

    headers = ["#", "Module", "Sheet", "Total", "Pass", "Fail", "Blocked", "Not Executed", "% Executed", "% Pass"]
    hr = 4
    for ci, h in enumerate(headers, 1):
        c = dash.cell(row=hr, column=ci, value=h)
        c.font = Font(bold=True, color=HEADER_FG, size=10)
        c.fill = PatternFill("solid", fgColor=HEADER_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    widths = [5, 26, 20, 9, 9, 9, 10, 14, 12, 11]
    for ci, w in enumerate(widths, 1):
        dash.column_dimensions[get_column_letter(ci)].width = w

    for i, m in enumerate(modules):
        r = hr + 1 + i
        sheet = m["sheet"]
        last = m["last"]
        rng = f"'{sheet}'!{STATUS_LTR}2:{STATUS_LTR}{last}"
        idrng = f"'{sheet}'!A2:A{last}"
        vals = [
            i + 1,
            m["name"],
            sheet,
            f"=COUNTA({idrng})",
            f'=COUNTIF({rng},"Pass")',
            f'=COUNTIF({rng},"Fail")',
            f'=COUNTIF({rng},"Blocked")',
            f'=COUNTIF({rng},"Not Executed")',
            f'=IFERROR((D{r}-H{r})/D{r},0)',
            f'=IFERROR(E{r}/(D{r}-H{r}),0)',
        ]
        for ci, v in enumerate(vals, 1):
            c = dash.cell(row=r, column=ci, value=v)
            c.border = BORDER
            c.font = Font(size=10, bold=(ci == 2))
            c.alignment = Alignment(horizontal="left" if ci == 2 else "center", vertical="center")
            if ci in (9, 10):
                c.number_format = "0.0%"
    # totals row
    tr = hr + 1 + len(modules)
    dash.cell(row=tr, column=2, value="TOTAL").font = Font(bold=True, size=10)
    for ci in range(4, 9):
        L = get_column_letter(ci)
        c = dash.cell(row=tr, column=ci, value=f"=SUM({L}{hr+1}:{L}{tr-1})")
        c.font = Font(bold=True, size=10); c.border = BORDER
        c.alignment = Alignment(horizontal="center")
    c = dash.cell(row=tr, column=9, value=f"=IFERROR((D{tr}-H{tr})/D{tr},0)")
    c.number_format = "0.0%"; c.font = Font(bold=True); c.border = BORDER; c.alignment = Alignment(horizontal="center")
    c = dash.cell(row=tr, column=10, value=f"=IFERROR(E{tr}/(D{tr}-H{tr}),0)")
    c.number_format = "0.0%"; c.font = Font(bold=True); c.border = BORDER; c.alignment = Alignment(horizontal="center")
    for ci in (1, 2, 3):
        dash.cell(row=tr, column=ci).border = BORDER
    dash.freeze_panes = "A5"

    # per-module sheets
    for m in modules:
        ws = wb.create_sheet(m["sheet"])
        write_test_sheet(ws, m["data"])

    path = os.path.join(OUT, "00_Saranam_Master_Test_Plan.xlsx")
    wb.save(path)
    return path


def main():
    files = sorted(glob.glob(os.path.join(BASE, "*", "*_test_cases.json")))
    modules = []
    for f in files:
        data = json.load(open(f))
        dirname = os.path.basename(os.path.dirname(f))          # e.g. 01-patient
        num = dirname.split("-")[0]
        safe = data["module"].replace("/", "-").replace(":", "")[:28]
        fname = f"{num}_{data['module'].replace(' ', '_').replace('/', '-')}_Test_Cases.xlsx"
        path, last = build_module_workbook(data, fname)
        modules.append({"code": data.get("module_code"), "name": data["module"],
                        "sheet": safe, "data": data, "fname": fname, "last": last})
        print(f"  built {fname:<48} ({len(data['test_cases'])} cases)")
    master = build_master(modules)
    print(f"  built {os.path.basename(master):<48} (dashboard + {len(modules)} module sheets)")
    print(f"\nTotal cases: {sum(len(m['data']['test_cases']) for m in modules)}")
    print(f"Output dir : {OUT}")


if __name__ == "__main__":
    main()
