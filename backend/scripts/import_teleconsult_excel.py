#!/usr/bin/env python3
"""
Import teleconsult visits and documentation from an Excel (.xlsx) file.

Expected Columns:
- Date of Contact (e.g. '12/03/2026' or '2026-03-12')
- Hospital UHID (e.g. '20210017422')
- Homecare UHID code (e.g. 'KAGBAKMIO')
- Home care UHID (e.g. '2026010001')
- First Name
- Last Name
- Type of call (e.g. 'Telephonic call', 'Video call')
- Type of consultation (e.g. 'Nurse', 'Doctor', 'Counsellor', 'Physiotherapist')
- Incoming/Outgoing call (e.g. 'Incoming call', 'Outgoing call')
- Call type (e.g. 'Follow-up', 'Condolence', 'Bereavement')
- Signs and symptoms
- Advice given

Usage:
    python scripts/import_teleconsult_excel.py /path/to/file.xlsx
    python scripts/import_teleconsult_excel.py /path/to/file.xlsx --dry-run
    python scripts/import_teleconsult_excel.py /path/to/file.xlsx --sheet "Phonecall service"
"""
import sys
import os
import argparse
from datetime import datetime, date
from pathlib import Path

# Add backend root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import create_app
from app.extensions import db
from app.patients.models import Patient
from app.saranam.models.patient_visit import PatientVisit
from app.teleconsult.models.tele_session import TeleSession, STATUS_COMPLETED
from app.teleconsult.models.teleconsult_documentation import TeleconsultDocumentation


def parse_date_value(val):
    if not val:
        return date.today()
    if isinstance(val, date):
        return val
    if isinstance(val, datetime):
        return val.date()

    val_str = str(val).strip()
    if not val_str:
        return date.today()

    for fmt in ["%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y", "%d-%m-%Y", "%d.%m.%Y"]:
        try:
            return datetime.strptime(val_str, fmt).date()
        except ValueError:
            pass

    return date.today()


def clean_str(val):
    if val is None:
        return None
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s if s else None


def normalize_header(h):
    if not h:
        return ""
    return str(h).strip().lower().replace("\n", " ").replace("_", " ")


def find_column_idx(header_map, possible_names):
    # 1. Try exact match first
    for name in possible_names:
        norm = normalize_header(name)
        for h_text, idx in header_map.items():
            if norm == h_text:
                return idx

    # 2. Try substring match
    for name in possible_names:
        norm = normalize_header(name)
        for h_text, idx in header_map.items():
            if norm in h_text:
                return idx
    return None


def main():
    parser = argparse.ArgumentParser(description="Import teleconsult records from Excel file")
    parser.add_argument("excel_path", help="Path to Excel (.xlsx) file")
    parser.add_argument("--sheet", help="Name of sheet to import (default: active sheet)", default=None)
    parser.add_argument("--dry-run", action="store_true", help="Simulate import without committing changes")
    args = parser.parse_args()

    excel_file = Path(args.excel_path)
    if not excel_file.exists():
        print(f"Error: File not found at {excel_file}")
        sys.exit(1)

    try:
        import openpyxl
    except ImportError:
        print("Error: openpyxl library is required. Run 'pip install openpyxl'")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_file, data_only=True)
    if args.sheet:
        if args.sheet not in wb.sheetnames:
            print(f"Error: Sheet '{args.sheet}' not found. Available sheets: {wb.sheetnames}")
            sys.exit(1)
        ws = wb[args.sheet]
    else:
        ws = wb.active

    print(f"Loading sheet: '{ws.title}' ({ws.max_row} rows)...")

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("Error: Sheet is empty.")
        sys.exit(1)

    # Build header map from row 1
    raw_headers = rows[0]
    header_map = {}
    for idx, h in enumerate(raw_headers):
        if h is not None:
            clean_h = " ".join(str(h).split()).lower()
            header_map[clean_h] = idx

    col_date = find_column_idx(header_map, ["date of contact", "contact date", "date"])
    col_hosp_uhid = find_column_idx(header_map, ["hospital uhid"])
    col_hc_code = find_column_idx(header_map, ["homecare uhid code", "home care uhid code", "homecare code"])
    col_hc_uhid = find_column_idx(header_map, ["home care uhid", "homecare uhid"])
    col_fname = find_column_idx(header_map, ["first name"])
    col_lname = find_column_idx(header_map, ["last name"])
    col_call_type_channel = find_column_idx(header_map, ["type of call"])
    col_consultation_type = find_column_idx(header_map, ["type of consultation"])
    col_call_direction = find_column_idx(header_map, ["incoming/outgoing call", "incoming/outgoing"])
    col_call_type_category = find_column_idx(header_map, ["call type"])
    col_symptoms = find_column_idx(header_map, ["signs and symptoms", "symptoms"])
    col_advice = find_column_idx(header_map, ["advice given", "advice"])

    print("Column Mapping Identified:")
    print(f"  - Date of Contact        : Col {col_date}")
    print(f"  - Hospital UHID          : Col {col_hosp_uhid}")
    print(f"  - Homecare UHID Code     : Col {col_hc_code}")
    print(f"  - Homecare UHID          : Col {col_hc_uhid}")
    print(f"  - First Name / Last Name : Col {col_fname} / Col {col_lname}")
    print(f"  - Type of Call (Channel) : Col {col_call_type_channel}")
    print(f"  - Type of Consultation   : Col {col_consultation_type}")
    print(f"  - Call Direction         : Col {col_call_direction}")
    print(f"  - Call Type Category     : Col {col_call_type_category}")
    print(f"  - Signs & Symptoms       : Col {col_symptoms}")
    print(f"  - Advice Given           : Col {col_advice}")
    print("-" * 60)

    app = create_app()
    with app.app_context():
        imported_count = 0
        skipped_count = 0

        for row_idx, row in enumerate(rows[1:], start=2):
            if not any(row):
                continue

            date_contact = parse_date_value(row[col_date]) if col_date is not None else date.today()
            hosp_uhid = clean_str(row[col_hosp_uhid]) if col_hosp_uhid is not None else None
            hc_code = clean_str(row[col_hc_code]) if col_hc_code is not None else None
            hc_uhid = clean_str(row[col_hc_uhid]) if col_hc_uhid is not None else None
            fname = clean_str(row[col_fname]) if col_fname is not None else ""
            lname = clean_str(row[col_lname]) if col_lname is not None else ""

            full_name = f"{fname or ''} {lname or ''}".strip() or "Unnamed Patient"

            raw_channel = clean_str(row[col_call_type_channel]) if col_call_type_channel is not None else "voice"
            channel = "video" if raw_channel and "video" in raw_channel.lower() else "voice"

            consultation_type = clean_str(row[col_consultation_type]) if col_consultation_type is not None else "Nurse"
            call_direction = clean_str(row[col_call_direction]) if col_call_direction is not None else "Outgoing call"
            call_type = clean_str(row[col_call_type_category]) if col_call_type_category is not None else "Follow-up"
            symptoms = clean_str(row[col_symptoms]) if col_symptoms is not None else None
            advice = clean_str(row[col_advice]) if col_advice is not None else None

            # Find Existing Patient Only (do NOT create new patient)
            patient = None
            if hc_uhid:
                patient = Patient.query.filter_by(uhid=hc_uhid).first()
            if not patient and hosp_uhid:
                patient = Patient.query.filter_by(hospital_uhid=hosp_uhid).first()
            if not patient and full_name and full_name != "Unnamed Patient":
                patient = Patient.query.filter(Patient.full_name.ilike(full_name)).first()

            if not patient:
                skipped_count += 1
                print(f"Row {row_idx}: SKIPPED — No existing patient found for '{full_name}' (HC UHID: {hc_uhid or 'N/A'}, Hosp UHID: {hosp_uhid or 'N/A'})")
                continue

            # Compute next visit_number for patient
            current_max = db.session.query(db.func.max(PatientVisit.visit_number))\
                .filter(PatientVisit.patient_id == patient.id).scalar() or 0
            visit_number = current_max + 1

            # Create PatientVisit
            visit = PatientVisit(
                patient_id=patient.id,
                visit_number=visit_number,
                visit_type="teleconsult",
                visit_date=date_contact,
                visit_location=channel,
                status="completed",
            )
            db.session.add(visit)
            db.session.flush()

            # Create TeleSession
            session = TeleSession(
                patient_id=patient.id,
                linked_visit_id=visit.id,
                channel=channel,
                status=STATUS_COMPLETED,
            )
            db.session.add(session)
            db.session.flush()

            # Create TeleconsultDocumentation
            doc = TeleconsultDocumentation(
                session_id=session.id,
                who_answered="patient",
                caller_name=patient.full_name,
                caller_phone=None,
                phone_used_type="person_phone",
                consultation_type=consultation_type,
                call_direction=call_direction,
                call_type=call_type,
                symptoms=symptoms,
                advice=advice,
            )
            db.session.add(doc)
            imported_count += 1

            print(f"Row {row_idx}: Imported Visit #{visit_number} for '{patient.full_name}' (UHID: {patient.uhid or 'N/A'}) - {date_contact}")

        if args.dry_run:
            db.session.rollback()
            print("\n[DRY RUN COMPLETE] No changes were committed to the database.")
            print(f"  - Teleconsult Visits Matched & Processed: {imported_count}")
            print(f"  - Skipped Rows (No Patient Match): {skipped_count}")
        else:
            db.session.commit()
            print("\n[IMPORT COMPLETE] All teleconsult records committed successfully!")
            print(f"  - Teleconsult Visits Imported: {imported_count}")
            print(f"  - Skipped Rows (No Patient Match): {skipped_count}")


if __name__ == "__main__":
    main()
