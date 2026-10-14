"""Submissions services."""
from datetime import datetime
from app.extensions import db
from app.submissions.models import FormSubmission
from app.forms.models import Form
from app.auth.decorators import get_current_user

# Backward compatibility alias
Submission = FormSubmission


class SubmissionService:
    """Submission management service."""

    @staticmethod
    def create_submission(form_id, data):
        """Create a new submission."""
        # Get form
        form = Form.query.get(form_id)
        if not form or not form.is_active:
            return None, "Form not found or inactive"

        # Create submission
        submission = Submission(
            form_id=form_id,
            data=data
        )
        submission.form = form  # Set form relationship for validation

        # Set current user
        current_user = get_current_user()
        if current_user:
            submission.submitted_by = current_user.id
        else:
            return None, "User not authenticated"

        # Validate submission
        is_valid, validation_result = submission.validate_against_form()
        if not is_valid:
            return None, {'message': 'Validation failed', 'errors': validation_result}

        try:
            db.session.add(submission)
            db.session.commit()
            return submission, "Submission created successfully"
        except Exception as e:
            db.session.rollback()
            return None, str(e)

    @staticmethod
    def get_submissions(page=1, per_page=50, form_id=None, user_id=None,
                        date_from=None, date_to=None):
        """Get paginated list of submissions with filters."""
        query = Submission.query

        # Apply filters
        if form_id:
            query = query.filter_by(form_id=form_id)
        if user_id:
            query = query.filter_by(submitted_by=user_id)
        if date_from:
            query = query.filter(Submission.created_at >= date_from)
        if date_to:
            query = query.filter(Submission.created_at <= date_to)

        # Get current user for filtering
        current_user = get_current_user()
        if current_user and current_user.role == 'data_entry':
            # Data entry users can only see their own submissions
            query = query.filter_by(submitted_by=current_user.id)

        paginated = query.order_by(Submission.created_at.desc()).paginate(
            page=page,
            per_page=per_page,
            error_out=False
        )

        return {
            'submissions': [sub.to_dict() for sub in paginated.items],
            'total': paginated.total,
            'pages': paginated.pages,
            'current_page': page
        }

    @staticmethod
    def get_submission_by_id(submission_id):
        """Get submission by ID."""
        submission = Submission.query.get(submission_id)
        if not submission:
            return None, "Submission not found"

        # Check permissions
        current_user = get_current_user()
        if current_user and current_user.role == 'data_entry':
            if submission.submitted_by != current_user.id:
                return None, "Access denied"

        return submission, "Submission found"

    @staticmethod
    def update_submission(submission_id, data):
        """Update existing submission."""
        from app.submissions.models import FormAnswer

        submission = Submission.query.get(submission_id)
        if not submission:
            return None, "Submission not found"

        # Check permissions
        current_user = get_current_user()
        if current_user:
            if current_user.role == 'data_entry' and submission.submitted_by != current_user.id:
                return None, "Access denied"
        else:
            return None, "User not authenticated"

        try:
            # Update data_backup for backward compatibility
            submission.data_backup = data
            submission.updated_at = datetime.utcnow()

            # Delete existing answers
            FormAnswer.query.filter_by(submission_id=submission.id).delete()

            # Create new answers from data
            for field_id, value in data.items():
                if value is not None and value != '':
                    answer = FormAnswer(
                        submission_id=submission.id,
                        field_id=field_id,
                        field_type='text',  # Auto-detected in set_value
                    )
                    answer.set_value(value)
                    db.session.add(answer)

            db.session.commit()
            return submission, "Submission updated successfully"
        except Exception as e:
            db.session.rollback()
            return None, str(e)

    @staticmethod
    def delete_submission(submission_id):
        """Delete a submission."""
        submission = Submission.query.get(submission_id)
        if not submission:
            return False, "Submission not found"

        # Check permissions (only admin can delete)
        current_user = get_current_user()
        if not current_user or current_user.role != 'admin':
            return False, "Access denied"

        try:
            db.session.delete(submission)
            db.session.commit()
            return True, "Submission deleted successfully"
        except Exception as e:
            db.session.rollback()
            return False, str(e)

    @staticmethod
    def export_submissions(form_id=None, format='json'):
        """Export submissions data in JSON, CSV, or Excel format."""
        query = Submission.query
        if form_id:
            query = query.filter_by(form_id=form_id)

        submissions = query.all()

        if not submissions:
            if format == 'json':
                return []
            elif format == 'csv':
                return ''
            elif format == 'xlsx':
                import io
                from openpyxl import Workbook
                wb = Workbook()
                ws = wb.active
                ws.title = "Submissions"
                ws.append(['No submissions found'])
                output = io.BytesIO()
                wb.save(output)
                output.seek(0)
                return output.getvalue()

        if format == 'json':
            return [sub.to_dict() for sub in submissions]

        # Get form schema to extract proper column headers
        form = submissions[0].form if submissions else None
        field_columns = []
        field_labels = {}

        if form and form.schema:
            schema = form.schema
            # Support both single-page and multi-page formats
            fields = schema.get('fields', [])
            if not fields and 'pages' in schema:
                fields = []
                for page in schema['pages']:
                    fields.extend(page.get('fields', []))

            # Build column mapping (field_id -> label)
            for field in fields:
                if field.get('type') != 'section':  # Exclude section headers
                    field_id = field.get('id')
                    field_label = field.get('label', field_id)
                    field_columns.append(field_id)
                    field_labels[field_id] = field_label
        else:
            # Fallback: get all unique field keys from submission data
            all_keys = set()
            for sub in submissions:
                if sub.data:
                    all_keys.update(sub.data.keys())
            field_columns = sorted(list(all_keys))
            field_labels = {k: k for k in field_columns}

        # Define standard columns
        standard_columns = ['id', 'form_name', 'submitted_by', 'created_at', 'status']
        all_column_ids = standard_columns + field_columns

        # For CSV format
        if format == 'csv':
            import csv
            import io

            output = io.StringIO()

            # Create headers with proper labels
            headers = ['ID', 'Form Name', 'Submitted By', 'Submitted At', 'Status']
            headers.extend([field_labels.get(fid, fid) for fid in field_columns])

            writer = csv.writer(output)
            writer.writerow(headers)

            for sub in submissions:
                row = [
                    str(sub.id),
                    sub.form.name if sub.form else '',
                    sub.submitted_by_user.full_name if sub.submitted_by_user else '',
                    sub.created_at.isoformat(),
                    sub.status
                ]

                # Add field values in correct order
                for field_id in field_columns:
                    value = sub.data.get(field_id, '') if sub.data else ''
                    row.append(str(value) if value is not None else '')

                writer.writerow(row)

            return output.getvalue()

        # For Excel format
        if format == 'xlsx':
            import io
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment

            wb = Workbook()
            ws = wb.active
            ws.title = "Submissions"

            # Create headers with proper labels
            headers = ['ID', 'Form Name', 'Submitted By', 'Submitted At', 'Status']
            headers.extend([field_labels.get(fid, fid) for fid in field_columns])

            ws.append(headers)

            # Style header row
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Add data rows
            for sub in submissions:
                row = [
                    str(sub.id),
                    sub.form.name if sub.form else '',
                    sub.submitted_by_user.full_name if sub.submitted_by_user else '',
                    sub.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                    sub.status
                ]

                # Add field values in correct order
                for field_id in field_columns:
                    value = sub.data.get(field_id, '') if sub.data else ''
                    row.append(value if value is not None else '')

                ws.append(row)

            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            # Save to BytesIO
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output.getvalue()

        return submissions